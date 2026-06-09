#!/usr/bin/env python3
"""
DB1 Report Transform — skill: db1-report-transform
Transforms raw DB1 task spreadsheet into the standardized DB1 report Excel format.
"""

import sys
import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Serviço mapping ───────────────────────────────────────────────────────────
def get_servico(tipo_tarefa, responsavel):
    try:
        tipo = int(pd.to_numeric(tipo_tarefa, errors='coerce'))
    except (TypeError, ValueError):
        return ''

    resp = str(responsavel).strip().upper() if pd.notna(responsavel) else ''

    if tipo == 15:  return 'TREINAMENTO'
    if tipo in (49, 67, 70): return 'DESENVOLVIMENTO'
    if tipo == 55:  return 'GESTÃO'
    if tipo == 56:  return 'REUNIÃO'
    if tipo == 75:
        if resp == 'GABIELA.COUTINHO': return 'UX'
        if resp == 'PRISCILA.SANTANA': return 'ANÁLISE'
        return 'REUNIÃO'
    return ''


def get_solicitante(cod):
    if pd.isna(cod): return ''
    try:
        cod = int(float(str(cod).strip()))
    except (ValueError, TypeError):
        return ''
    if cod == 667: return 'SUPRIMENTOS'
    if cod == 666: return 'COMERCIAL'
    return ''


def get_project_id(cod):
    if pd.isna(cod): return ''
    try:
        cod = int(float(str(cod).strip()))
    except (ValueError, TypeError):
        return ''
    if cod == 667: return 'Portal'
    if cod == 666: return 'Ventas'
    return ''


def _norm(s):
    """Lowercase + strip accents for fuzzy column matching."""
    import unicodedata
    return unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().strip().lower()

def find_col(df, candidates):
    """Find column by partial/normalized name."""
    norm = {_norm(c): c for c in df.columns}
    for cand in candidates:
        k = _norm(cand)
        if k in norm:
            return norm[k]
        # partial match
        for col_key, col_orig in norm.items():
            if k in col_key or col_key in k:
                return col_orig
    return None


# ── Transform ─────────────────────────────────────────────────────────────────
def transform_db1(df):
    df.columns = [str(c).strip() for c in df.columns]

    # Locate source columns flexibly
    col_equipe   = find_col(df, ['Cód. Equipe Responsável', 'Cod. Equipe Responsavel', 'Equipe'])
    col_resp     = find_col(df, ['Cód. Responsavel', 'Cod. Responsavel', 'Responsavel'])
    col_tipo     = find_col(df, ['Cód. Tipo Tarefa', 'Cod. Tipo Tarefa', 'Tipo Tarefa'])
    col_tarefa   = find_col(df, ['Cód. Tarefa', 'Cod. Tarefa', 'Tarefa'])
    col_desc     = find_col(df, ['Desc Tarefa', 'Descrição', 'Descricao', 'Descricao da Tarefa'])
    col_inicio   = find_col(df, ['Dh início', 'Dh inicio', 'Data Inicio', 'Inicio'])
    col_termino  = find_col(df, ['Dh Termino', 'Dh Término', 'Data Termino', 'Termino'])
    col_total    = find_col(df, ['Total', 'TOTAL', 'Horas', 'Qtd Horas'])
    col_inf      = find_col(df, ['Inf Compl Execucao', 'Inf Compl', 'Observacao', 'Comentário'])

    def safe_series(col):
        return df[col] if col and col in df.columns else pd.Series([''] * len(df), index=df.index)

    out = pd.DataFrame(index=df.index)
    out['Empresa']   = 'DB1'
    out['Solicitante'] = safe_series(col_equipe).apply(get_solicitante)
    out['Serviço']   = df.apply(
        lambda r: get_servico(r.get(col_tipo, ''), r.get(col_resp, '')), axis=1
    ) if col_tipo else pd.Series([''] * len(df))

    out['Prioridade'] = 'NORMAL'
    out['Descrição']  = safe_series(col_desc)
    out['Usuário']    = ''

    dh_inicio = pd.to_datetime(safe_series(col_inicio), dayfirst=True, errors='coerce')
    out['Document Date']   = dh_inicio.dt.strftime('%d/%m/%Y').fillna('')
    out['Start Date/Time'] = dh_inicio.dt.strftime('%d/%m/%Y %H:%M').fillna('')

    dh_termino = pd.to_datetime(safe_series(col_termino), dayfirst=True, errors='coerce')
    out['End Date/Time']   = dh_termino.dt.strftime('%d/%m/%Y %H:%M').fillna('')

    out['Quantity']          = pd.to_numeric(safe_series(col_total), errors='coerce').fillna(0)
    out['Price']             = 229.92
    out['Amount']            = '__FORMULA__'   # replaced by Excel formula in write step
    out['Currency']          = 'Real'
    out['Travel Time']       = ''
    out['Bookable Resource'] = safe_series(col_resp)
    out['Resource role ID']  = ''
    out['Project ID']        = safe_series(col_equipe).apply(get_project_id)
    out['Task ID']           = safe_series(col_tarefa)
    out['Observação']        = safe_series(col_inf)

    return out


# ── Excel writer ──────────────────────────────────────────────────────────────
def write_report(out_df, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório DB1"

    headers = list(out_df.columns)

    header_fill = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
    header_font = Font(name='Arial', bold=True, color="FFFFFF", size=11)
    center      = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin        = Side(style='thin', color='BFBFBF')
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    row_font    = Font(name='Arial', size=10)

    price_col  = headers.index('Price')  + 1
    qty_col    = headers.index('Quantity') + 1
    amount_col = headers.index('Amount')   + 1

    # Header row
    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = border
    ws.row_dimensions[1].height = 30

    # Data rows
    for ri, row in enumerate(out_df.itertuples(index=False), start=2):
        for ci, value in enumerate(row, start=1):
            col_name = headers[ci - 1]
            if col_name == 'Amount':
                p_ref = f"{get_column_letter(price_col)}{ri}"
                q_ref = f"{get_column_letter(qty_col)}{ri}"
                cell  = ws.cell(row=ri, column=ci, value=f"={p_ref}*{q_ref}")
            else:
                if isinstance(value, float) and np.isnan(value):
                    value = ''
                cell = ws.cell(row=ri, column=ci, value=value)
            cell.font      = row_font
            cell.alignment = Alignment(vertical='center')
            cell.border    = border

        # Zebra striping
        if ri % 2 == 0:
            fill = PatternFill("solid", start_color="EBF3FB", end_color="EBF3FB")
            for ci in range(1, len(headers) + 1):
                ws.cell(row=ri, column=ci).fill = fill

    # Column widths
    for ci, h in enumerate(headers, start=1):
        col_letter = get_column_letter(ci)
        max_len    = max(len(str(h)), 12)
        for row_cells in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
            for cell in row_cells:
                try:
                    max_len = max(max_len, len(str(cell.value or '')))
                except Exception:
                    pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    ws.freeze_panes = 'A2'
    wb.save(output_path)
    print(f"[OK] Relatório salvo: {output_path}")


# ── Main ──────────────────────────────────────────────────────────────────────
def find_header_row(filepath, sheet_name=0, max_rows=30):
    """Detect the row index that contains the real column headers."""
    probe = pd.read_excel(filepath, sheet_name=sheet_name, header=None, nrows=max_rows)
    keywords = ['desc tarefa', 'dh inicio', 'dh início', 'total', 'cod. responsavel',
                'cód. responsavel', 'cód. equipe', 'cod. equipe', 'inf compl']
    for i, row in probe.iterrows():
        vals = [str(v).strip().lower() for v in row.values if pd.notna(v)]
        if any(any(k in v for k in keywords) for v in vals):
            print(f"[INFO] Cabeçalho detectado na linha {i}")
            return i
    return 0

def run(input_path: str, output_path: str):
    print(f"[INFO] Lendo: {input_path}")
    sheet_names = list(pd.read_excel(input_path, sheet_name=None, header=None, nrows=1).keys())

    target_df = None
    for name in sheet_names:
        header_row = find_header_row(input_path, sheet_name=name)
        sdf = pd.read_excel(input_path, sheet_name=name, header=header_row)
        sdf.columns = [str(c).strip() for c in sdf.columns]
        cols_lower = [c.lower() for c in sdf.columns]
        if any('desc tarefa' in c or 'dh in' in c or 'inf compl' in c for c in cols_lower):
            target_df = sdf
            print(f"[INFO] Sheet selecionada: '{name}' | header row: {header_row}")
            break

    if target_df is None:
        name = sheet_names[0]
        header_row = find_header_row(input_path, sheet_name=name)
        target_df = pd.read_excel(input_path, sheet_name=name, header=header_row)
        target_df.columns = [str(c).strip() for c in target_df.columns]
        print("[WARN] Usando primeira sheet com header detectado.")

    out_df = transform_db1(target_df)
    write_report(out_df, output_path)
    print(f"[OK] {len(out_df)} linhas transformadas.")


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('--input',  required=True)
    ap.add_argument('--output', default='relatorio_db1.xlsx')
    args = ap.parse_args()
    run(args.input, args.output)
